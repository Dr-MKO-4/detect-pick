"""Page Données liste des fichiers XLSX + prévisualisation réelle."""
import glob as _glob
import os
import logging
from dash import html, dcc
from config import XLSX_PATTERN, PAYS, VOLETS
from layout.icons import svg, ICO_FILE, ICO_FOLDER, ICO_GRID, ICO_DL, ICO_CHECK, ICO_WARN, ICO_TRASH

log = logging.getLogger(__name__)

# Cache : (path, mtime) → meta dict  (invalidé automatiquement si le fichier change)
_meta_cache: dict = {}

# Cache : (path, mtime) → DataFrame complet (invalidé si le fichier change)
_full_df_cache: dict = {}


def _read_file_full(path: str):
    """Lit le fichier XLSX en entier (toutes lignes), mis en cache par (path, mtime).

    Utilisé pour l'export CSV — distinct de `_read_file_meta` dont le DataFrame
    est tronqué à 200 lignes pour la prévisualisation.
    """
    try:
        mtime = os.path.getmtime(path)
    except OSError:
        mtime = 0

    cache_key = (path, mtime)
    if cache_key in _full_df_cache:
        return _full_df_cache[cache_key]

    import pandas as pd
    df = pd.read_excel(path, index_col=0)
    _full_df_cache[cache_key] = df
    return df


def _first_date_col(cols: list) -> str:
    """Retourne le premier header qui ressemble à une période (date, mois, trimestre…)."""
    import re
    _date_re = re.compile(
        r'(\d{4}[Mm]\d{2})'          # 2000M01
        r'|(\d{4}[Qq]\d)'             # 2000Q1
        r'|([A-Za-z]{3}[-\s]\d{4})'  # Jan 2000
        r'|(\d{2}/\d{4})'             # 01/2000
        r'|(\d{4}-\d{2})',            # 2000-01
    )
    for c in cols:
        if _date_re.search(str(c)):
            return str(c)
    # Fallback : première colonne numérique (année seule ex. 2000)
    for c in cols:
        try:
            v = int(str(c))
            if 1900 <= v <= 2100:
                return str(c)
        except Exception:
            pass
    return cols[0] if cols else "—"


def _read_file_meta(path: str) -> dict:
    """Lit les métadonnées réelles du fichier XLSX — résultat mis en cache par (path, mtime)."""
    try:
        mtime = os.path.getmtime(path)
    except OSError:
        mtime = 0

    cache_key = (path, mtime)
    if cache_key in _meta_cache:
        return _meta_cache[cache_key]

    try:
        import pandas as pd
        df = pd.read_excel(path, index_col=0, nrows=200)
        # Exclure les colonnes non-numériques (ex. "IFS Code") du comptage des périodes
        num_cols = df.select_dtypes(include="number").columns.tolist()
        df_num   = df[num_cols] if num_cols else df
        n_ind    = df_num.shape[0]
        n_obs    = df_num.shape[1]
        pct_na   = round(df_num.isnull().values.mean() * 100, 1) if n_obs else "?"
        cols_str = [str(c) for c in df_num.columns]
        debut    = _first_date_col(cols_str) if cols_str else "—"
        fin      = cols_str[-1] if cols_str else "—"
        meta = {"n_ind": n_ind, "n_obs": n_obs, "pct_na": pct_na,
                "debut": debut, "fin": fin, "df": df_num, "ok": True}
    except Exception as exc:
        log.warning("Lecture XLSX impossible : %s", exc)
        meta = {"n_ind": "?", "n_obs": "?", "pct_na": "?",
                "debut": "—", "fin": "—", "df": None, "ok": False}

    _meta_cache[cache_key] = meta
    return meta


def _read_file_meta_light(path: str) -> dict:
    """Lit uniquement n_obs / n_ind sans charger tout le contenu (rapide).

    Une seule ouverture du fichier en mode `read_only` (openpyxl) pour ne
    lire que les dimensions de la feuille, au lieu de deux appels
    `pd.read_excel` séparés (un pour l'en-tête, un pour l'index).
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            n_obs = max(0, (ws.max_column or 1) - 1)
            n_ind = max(0, (ws.max_row or 1) - 1)
        finally:
            wb.close()
        return {"n_obs": n_obs, "n_ind": n_ind, "ok": True}
    except Exception:
        return {"n_obs": "?", "n_ind": "?", "ok": True}


def _preview_table_from_df(df) -> html.Table:
    """Génère un tableau de prévisualisation depuis un DataFrame réel."""
    if df is None:
        return html.Div("Impossible de lire le fichier.",
                        style={"padding": "14px", "fontSize": "11px", "color": "var(--red)"})

    cols_all = list(df.columns)
    display_cols = cols_all[:5] + [cols_all[-1]] if len(cols_all) > 5 else cols_all

    head_cells = [html.Th("Indicateur",
                           style={"position": "sticky", "left": "0",
                                  "background": "var(--surf)", "zIndex": "1"})]
    for c in display_cols[:5]:
        head_cells.append(html.Th(str(c)[:8]))
    if len(cols_all) > 5:
        head_cells.append(html.Th("…", style={"color": "var(--muted2)"}))
        head_cells.append(html.Th(str(display_cols[-1])[:8]))

    body_rows = []
    for i, (idx, row) in enumerate(df.iterrows()):
        if i >= 6:
            body_rows.append(html.Tr([
                html.Td(f"… {df.shape[0] - 6} autres indicateurs",
                        colSpan=len(head_cells),
                        style={"color": "var(--muted2)", "textAlign": "center",
                               "fontSize": "10px", "fontStyle": "italic"}),
            ]))
            break
        cells = [html.Td(str(idx)[:30],
                         style={"fontWeight": "500", "position": "sticky", "left": "0",
                                "background": "var(--surf2)", "fontSize": "11px",
                                "maxWidth": "200px", "overflow": "hidden",
                                "textOverflow": "ellipsis", "whiteSpace": "nowrap"})]
        for c in display_cols[:5]:
            v = row.get(c)
            is_na = v is None or (hasattr(v, '__class__') and str(type(v)) == "<class 'float'>"
                                  and str(v) == 'nan')
            try:
                import math
                is_na = is_na or math.isnan(float(v))
            except Exception:
                is_na = False
            if is_na:
                cells.append(html.Td("NA", className="anomaly",
                                     style={"background": "var(--red-lt)", "color": "var(--red)",
                                            "textAlign": "center"}))
            else:
                try:
                    cells.append(html.Td(f"{float(v):,.1f}".replace(",", " "),
                                         className="normal"))
                except Exception:
                    cells.append(html.Td(str(v)[:10], className="normal"))
        if len(cols_all) > 5:
            cells.append(html.Td("…", style={"color": "var(--muted2)"}))
            v = row.get(display_cols[-1])
            try:
                cells.append(html.Td(f"{float(v):,.1f}".replace(",", " "), className="normal"))
            except Exception:
                cells.append(html.Td("", className="anomaly"))

        body_rows.append(html.Tr(cells))

    return html.Table(className="data-table", children=[
        html.Thead(html.Tr(head_cells)),
        html.Tbody(body_rows),
    ])


def _stats_panel(df) -> html.Div:
    """Tableau de statistiques descriptives (describe) pour le DataFrame sélectionné."""
    if df is None:
        return html.Div("Sélectionnez un fichier pour afficher les statistiques.",
                        style={"padding": "14px", "fontSize": "11px", "color": "var(--muted2)"})
    try:
        import pandas as pd
        desc = df.T.describe().T
        desc = desc.rename(columns={"count": "N", "mean": "Moy.", "std": "Écart-type",
                                     "min": "Min", "25%": "Q1", "50%": "Médiane",
                                     "75%": "Q3", "max": "Max"})
        show_cols = [c for c in ["N", "Moy.", "Écart-type", "Min", "Q1", "Médiane", "Q3", "Max"]
                     if c in desc.columns]
        display = desc[show_cols].head(20)

        head = [html.Th("Indicateur", style={"position": "sticky", "left": "0",
                                              "background": "var(--surf)", "zIndex": "1"})]
        for c in show_cols:
            head.append(html.Th(c, style={"textAlign": "right", "whiteSpace": "nowrap"}))

        rows = []
        for idx, row in display.iterrows():
            cells = [html.Td(str(idx)[:35],
                             style={"fontWeight": "500", "position": "sticky", "left": "0",
                                    "background": "var(--surf2)", "fontSize": "10px",
                                    "maxWidth": "200px", "overflow": "hidden",
                                    "textOverflow": "ellipsis", "whiteSpace": "nowrap"})]
            for c in show_cols:
                v = row.get(c)
                try:
                    txt = f"{float(v):,.2f}".replace(",", " ")
                except Exception:
                    txt = str(v) if v is not None else "—"
                cells.append(html.Td(txt, style={"textAlign": "right",
                                                  "fontFamily": "Roboto Mono, monospace",
                                                  "fontSize": "10px"}))
            rows.append(html.Tr(cells))

        if len(desc) > 20:
            rows.append(html.Tr([
                html.Td(f"… {len(desc) - 20} indicateurs supplémentaires",
                        colSpan=len(show_cols) + 1,
                        style={"color": "var(--muted2)", "textAlign": "center",
                               "fontSize": "10px", "fontStyle": "italic"})
            ]))

        return html.Table(className="data-table", children=[
            html.Thead(html.Tr(head)),
            html.Tbody(rows),
        ])
    except Exception as exc:
        return html.Div(f"Erreur statistiques : {exc}",
                        style={"padding": "14px", "fontSize": "11px", "color": "var(--red)"})


# Cache du dernier listing (évite de re-scanner data/ à chaque re-rendu de page
# déclenché par un store sans rapport, ex. polling de progression du pipeline).
_scan_cache: dict = {"files": None}


def scan_files_cached() -> list[dict]:
    """Retourne le dernier listing de fichiers, en le calculant si absent."""
    if _scan_cache["files"] is None:
        _scan_cache["files"] = scan_files()
    return _scan_cache["files"]


def refresh_scan_cache() -> list[dict]:
    """Force un nouveau scan et met à jour le cache (navigation vers Données)."""
    _scan_cache["files"] = scan_files()
    return _scan_cache["files"]


def invalidate_scan_cache() -> None:
    """À appeler après toute modification du dossier data/ (suppression, ajout)."""
    _scan_cache["files"] = None


def scan_files() -> list[dict]:
    """Scanne le dossier data/ et lit les métadonnées de tous les fichiers."""
    files = sorted(_glob.glob(XLSX_PATTERN))
    if not files:
        result = []
        for pays in PAYS:
            for volet in VOLETS:
                name = f"clean_{pays}_beac_mapping_2SR_{volet}.xlsx"
                result.append({"name": name, "meta": "Fichier manquant", "ok": False})
        return result

    result = []
    for i, f in enumerate(files):
        name = os.path.basename(f)
        if i == 0:
            meta = _read_file_meta(f)
        else:
            meta = _read_file_meta_light(f)
        desc = (f"{meta.get('n_obs','?')} obs · {meta.get('n_ind','?')} ind. · "
                f"{meta.get('debut','—')} – {meta.get('fin','—')}")
        entry = {"name": name, "ok": True, "path": f,
                 "meta": desc, "selected": i == 0}
        if i == 0:
            entry["_meta"] = meta
        result.append(entry)
    return result


def _stat_cards(meta: dict) -> list:
    """Génère les 4 cartes de stats pour le fichier sélectionné."""
    n_obs  = meta.get("n_obs",  "—")
    n_ind  = meta.get("n_ind",  "—")
    pct_na = meta.get("pct_na", "—")
    debut  = meta.get("debut",  "—")
    return [
        html.Div(className="stat-card", children=[
            html.Div(str(n_obs), className="stat-num"),
            html.Div("Périodes", className="stat-label"),
        ]),
        html.Div(className="stat-card", children=[
            html.Div(str(n_ind), className="stat-num"),
            html.Div("Indicateurs IFS", className="stat-label"),
        ]),
        html.Div(className="stat-card", children=[
            html.Div(f"{pct_na}%" if pct_na not in ("—", "?") else "—",
                     className="stat-num red"),
            html.Div("Valeurs manquantes", className="stat-label"),
        ]),
        html.Div(className="stat-card", children=[
            html.Div(str(debut),
                     className="stat-num",
                     style={"fontSize": "13px", "color": "var(--green)"}),
            html.Div("Début série", className="stat-label"),
        ]),
    ]


def _preview_panel(meta: dict | None) -> html.Div:
    df = meta.get("df") if meta else None
    n_ind = meta.get("n_ind", "?") if meta else "?"
    n_obs = meta.get("n_obs", "?") if meta else "?"

    return html.Div(className="panel", style={"margin": "0", "flex": "1"}, children=[
        html.Div(className="panel-header", children=[
            svg(ICO_GRID, size=12, stroke="var(--gold)"),
            html.Span("Prévisualisation  6 premières lignes", className="panel-title"),
            html.Div(style={"marginLeft": "auto", "display": "flex", "gap": "6px"}, children=[
                html.Button([svg(ICO_DL, size=11), " CSV"],
                            id="btn-export-csv", className="btn btn-ghost",
                            style={"padding": "4px 10px", "fontSize": "10px"}),
            ]),
        ]),
        html.Div(id="preview-table-container", style={"overflowX": "auto"},
                 children=[_preview_table_from_df(df)]),
        html.Div(style={"padding": "6px 12px", "borderTop": "1px solid var(--border)",
                        "display": "flex", "gap": "12px", "fontSize": "10px",
                        "color": "var(--muted2)"},
                 children=[
                     html.Span("Valeur manquante", style={"color": "var(--red)"}),
                     html.Span("Unité : Milliards FCFA"),
                     html.Span(f"Lignes : {n_ind} · Colonnes : {n_obs}",
                               style={"marginLeft": "auto"}),
                 ]),
        dcc.Download(id="dl-export-csv"),
    ])


def _build_file_rows(files_info: list, selected_idx: int = 0) -> list:
    """Génère la liste des divs de fichiers avec mise en évidence de la sélection."""
    rows = []
    for i, f in enumerate(files_info):
        is_sel = (i == selected_idx)
        icon_stroke = ("var(--gold)" if is_sel
                       else "var(--muted)" if not f.get("ok") else "var(--green)")
        badge = (html.Span(className="tag tag-green",
                           children=[svg(ICO_CHECK, size=8, stroke="currentColor")])
                 if f.get("ok")
                 else html.Span(className="tag tag-grey",
                                children=[svg(ICO_WARN, size=8, stroke="currentColor")]))
        del_btn = (html.Button(
            id={"type": "file-del", "index": i},
            n_clicks=0,
            title="Supprimer ce fichier",
            **{"aria-label": f"Supprimer le fichier {f['name']}"},
            className="btn btn-ghost",
            style={"padding": "2px 5px", "opacity": "0.45", "flexShrink": "0"},
            children=[svg(ICO_TRASH, size=11, stroke="var(--red)")],
        ) if f.get("ok") else None)

        rows.append(html.Div(
            id={"type": "file-row", "index": i},
            n_clicks=0,
            role="button", tabIndex=0,
            **({"aria-current": "true"} if is_sel else {}),
            style={"display": "flex", "alignItems": "center", "gap": "8px",
                   "padding": "7px 12px", "borderBottom": "1px solid var(--border)",
                   "cursor": "pointer",
                   "background": "var(--gold-lt)" if is_sel else "transparent"},
            children=[
                svg(ICO_FILE, size=12, stroke=icon_stroke),
                html.Div(style={"flex": "1", "minWidth": "0"}, children=[
                    html.Div(f["name"],
                             style={"fontSize": "11px",
                                    "fontWeight": "500" if is_sel else "400",
                                    "color": "var(--gold)" if is_sel else "var(--text)",
                                    "whiteSpace": "nowrap", "overflow": "hidden",
                                    "textOverflow": "ellipsis"}),
                    html.Div(f.get("meta", ""),
                             style={"fontSize": "9px", "color": "var(--muted)", "marginTop": "1px"}),
                ]),
                badge,
                del_btn,
            ],
        ))
    return rows


def donnees_layout(files_info: list | None = None) -> html.Div:
    files_info = files_info or scan_files()
    n_ok = len([f for f in files_info if f.get("ok")])

    selected = next((f for f in files_info if f.get("selected")), None)
    meta = (selected or {}).get("_meta") or {}
    n_obs  = meta.get("n_obs",  "—")
    n_ind  = meta.get("n_ind",  "—")
    pct_na = meta.get("pct_na", "—")
    debut  = meta.get("debut",  "—")
    df     = meta.get("df")

    file_rows = _build_file_rows(files_info, selected_idx=0)
    sel_name  = (selected or {}).get("name", "Aucun fichier sélectionné")
    n_total   = len(files_info)

    return html.Div(children=[
        html.Div(className="topbar", children=[
            html.Span("Données", className="tb-section"),
            html.Span("›", className="tb-sep"),
            html.Span("Fichiers source XLSX", className="tb-page"),
            html.Span(f"{n_ok} fichier{'s' if n_ok > 1 else ''} disponible{'s' if n_ok > 1 else ''}  ·  {n_total} chargé{'s' if n_total > 1 else ''}",
                      className="tb-context"),
        ]),
        html.Div(className="content", children=[
            html.Div(className="grid-donnees", children=[
                # Colonne gauche — liste fichiers
                html.Div(children=[
                    html.Div(
                        style={"border": "1.5px dashed rgba(158,111,26,.35)",
                               "background": "var(--gold-lt)",
                               "padding": "20px", "textAlign": "center",
                               "marginBottom": "14px", "cursor": "pointer"},
                        id="btn-open-folder", n_clicks=0,
                        role="button", tabIndex=0,
                        children=[
                            html.Div(svg(ICO_FOLDER, size=24, stroke="var(--gold-md)"),
                                     style={"marginBottom": "8px"}),
                            html.Div("Ouvrir dossier data/",
                                     style={"fontSize": "12px", "fontWeight": "600",
                                            "color": "var(--muted)"}),
                            html.Div("Déposer des fichiers XLSX ici",
                                     style={"fontSize": "10px", "color": "var(--muted2)",
                                            "marginTop": "4px"}),
                        ]),
                    html.Div(className="panel", style={"margin": "0"}, children=[
                        html.Div(className="panel-header", children=[
                            svg(ICO_FOLDER, size=12, stroke="var(--gold)"),
                            html.Span("Fichiers chargés", className="panel-title"),
                            html.Span(f"{n_ok} / {len(files_info)}", className="panel-sub"),
                        ]),
                        html.Div(id="file-list-container", children=file_rows),
                    ]),
                ]),

                # Colonne droite — info + stats + prévisualisation
                html.Div(style={"display": "flex", "flexDirection": "column", "gap": "14px"},
                         children=[
                    html.Div(
                        style={"display": "flex", "alignItems": "center", "gap": "10px",
                               "padding": "10px 14px", "background": "var(--surf)",
                               "border": "1px solid var(--border)"},
                        id="file-info-bar",
                        children=[
                            html.Span("Sélectionné", className="tag tag-gold"),
                            html.Span(sel_name,
                                      style={"fontSize": "12px", "fontWeight": "600",
                                             "overflow": "hidden", "textOverflow": "ellipsis",
                                             "whiteSpace": "nowrap", "maxWidth": "300px"}),
                            html.Span(f"{n_obs} périodes · {n_ind} indicateurs",
                                      style={"marginLeft": "auto", "fontSize": "10px",
                                             "color": "var(--muted)", "whiteSpace": "nowrap"}),
                        ]),
                    html.Div(id="donnees-stat-cards", className="grid-4",
                             children=_stat_cards(meta)),
                    _preview_panel(meta if meta else None),
                    # Statistiques descriptives
                    html.Div(className="panel", style={"margin": "0"}, children=[
                        html.Div(className="panel-header", children=[
                            svg(ICO_GRID, size=12, stroke="var(--gold)"),
                            html.Span("Statistiques descriptives", className="panel-title"),
                            html.Span("N · Moy. · Écart-type · Min · Q1 · Médiane · Q3 · Max",
                                      className="panel-sub"),
                        ]),
                        html.Div(id="donnees-stats-container",
                                 style={"overflowX": "auto"},
                                 children=[_stats_panel(df)]),
                    ]),
                ]),
            ]),
        ]),
    ])
